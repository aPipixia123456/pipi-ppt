"""Adapt Presenton's pinned v2 layout catalog to the common editable scene.

This deployment reuses upstream layout positions, text slots, colors and style
descriptions. All generation is routed through an explicit per-job gateway.
The upstream global-provider and binary export services are not exposed.
"""

import io
import json
import re
from copy import deepcopy
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .config import settings
from .schema import Element, Slide
from .storage import asset_path, store_asset

BUILTINS = {
    "executive": "工作汇报 / Reports",
    "modern": "解决方案 / Proposals",
    "momentum": "商业计划 / Business plans",
    "general": "教学课件 / Education",
    "nova": "产品介绍 / Product",
    "signal": "数据报告 / Data",
}


def color(value, default="#172033"):
    return (
        value.upper()
        if isinstance(value, str) and re.fullmatch(r"#[0-9a-fA-F]{6}", value)
        else default
    )


def upstream_elements(elements, x=0, y=0, template_id=""):
    for element in elements:
        position = element.get("position", {})
        left, top = x + position.get("x", 0), y + position.get("y", 0)
        if element["type"] == "group":
            yield from upstream_elements(element.get("children", []), left, top, template_id)
            continue
        size = element.get("size", {})
        width, height = size.get("width", 0), size.get("height", 0)
        if element["type"] == "vector":
            points = element.get("points", [])
            if not points:
                continue
            minimum_x, minimum_y = (
                min(p["x"] for p in points),
                min(p["y"] for p in points),
            )
            left, top = left + minimum_x, top + minimum_y
            width, height = (
                max(p["x"] for p in points) - minimum_x,
                max(p["y"] for p in points) - minimum_y,
            )
        if (
            width <= 0
            or height <= 0
            or left >= 1280
            or top >= 720
            or left + width <= 0
            or top + height <= 0
        ):
            continue
        original_box = dict(x=left, y=top, w=width, h=height)
        left, top = max(0, left), max(0, top)
        base = dict(
            id=element.get("name", "element")[:70],
            x=left,
            y=top,
            w=min(width, 1280 - left),
            h=min(height, 720 - top),
        )
        kind = element["type"]
        if kind == "text":
            font = element.get("font", {})
            text = "".join(run.get("text", "") for run in element.get("runs", [])) or element.get(
                "text", ""
            )
            yield Element(
                **base,
                type="text",
                text=text[:6000],
                color=color(font.get("color")),
                font_size=max(8, min(150, font.get("size", 24))),
                line_height=max(0.8, min(2.5, font.get("line_height", 1.25))),
                align=element.get("alignment", {}).get("horizontal", "left"),
                valign=element.get("alignment", {}).get("vertical", "top"),
                bold=bool(font.get("bold")),
                editable=not element.get("decorative", False),
                max_chars=max(1, min(6000, element.get("max_length", 300))),
            )
        elif kind == "image":
            # Template photos are content slots. No remote URL or tracking asset
            # from an upstream template is fetched by the public deployment.
            reference = f"{template_id}/{element.get('data', '')}"
            try:
                builtin_asset_path(reference)
            except ValueError:
                reference = None
            if (
                element.get("decorative")
                and width <= 3840
                and height <= 2160
                and left >= -2560
                and top >= -1440
            ):
                base.update(original_box)
            yield Element(
                **base,
                type="image",
                editable=not element.get("decorative", False),
                builtin_asset=reference,
                image_fit=element.get("fit", "contain")
                if element.get("fit") in {"contain", "fill", "cover"}
                else "contain",
            )
        elif kind == "chart":
            series = element.get("series", [{}])[0]
            labels = [str(item)[:100] for item in element.get("categories", [])][:30]
            values = [float(item or 0) for item in series.get("values", [])][:30]
            if len(labels) == len(values):
                yield Element(
                    **base,
                    type="chart",
                    labels=labels,
                    values=values,
                    fill=color(element.get("colors", ["#7058D2"])[0]),
                    color=color(element.get("text_color")),
                )
        elif kind == "vector":
            yield Element(
                **base,
                type="shape",
                editable=False,
                fill=color(element.get("fill", {}).get("color"), "#FFFFFF"),
                shape="ellipse" if element.get("shape") == "ellipse" else "rect",
            )


def builtin_template(template_id: str) -> dict:
    if template_id not in BUILTINS:
        raise ValueError("unknown_template")
    source = json.loads(
        (settings().template_dir / template_id / "template.json").read_text(encoding="utf-8")
    )
    layouts = []
    for raw in source["layouts"][:20]:
        elements = []
        for component in raw.get("components", []):
            position = component.get("position", {})
            elements.extend(
                upstream_elements(
                    component.get("elements", []),
                    position.get("x", 0),
                    position.get("y", 0),
                    template_id,
                )
            )
        for i, element in enumerate(elements):
            element.id = f"slot-{i}"
        if any(e.type == "text" and e.editable for e in elements):
            layouts.append(
                Slide(
                    id=raw["id"][:100],
                    name=raw.get("description", raw["id"])[:200],
                    elements=elements[:200],
                ).model_dump()
            )
    if not layouts:
        raise ValueError("template_has_no_editable_layouts")
    return {
        "id": template_id,
        "name": BUILTINS[template_id],
        "layouts": layouts,
        "theme": source.get("theme", {}),
        "warnings": [
            "Fonts are substituted with Noto Sans CJK SC. Complex vector paths are simplified; animations and SmartArt are not imported."
        ],
        "source": "Presenton 2b5078ba (Apache-2.0)",
    }


def builtin_asset_path(reference: str) -> Path:
    template_id, separator, relative = reference.partition("/")
    if not separator or template_id not in BUILTINS or not relative.startswith("static/"):
        raise ValueError("invalid_template_asset")
    root = (settings().template_dir / template_id / "static").resolve()
    path = (settings().template_dir / template_id / relative).resolve()
    if (
        not path.is_relative_to(root)
        or path.suffix.lower() not in {".png", ".jpg", ".jpeg", ".webp", ".svg"}
        or not path.is_file()
    ):
        raise ValueError("invalid_template_asset")
    return path


def pptx_color(rgb, fallback="#172033"):
    try:
        return color("#" + str(rgb.rgb), fallback)
    except (AttributeError, TypeError):
        return fallback


def imported_template(asset, job_id: str) -> dict:
    presentation = Presentation(asset_path(asset))
    if not 1 <= len(presentation.slides) <= 30:
        raise ValueError("template_requires_1_to_30_slides")
    scale = min(1280 / presentation.slide_width, 720 / presentation.slide_height)
    offset_x = (1280 - presentation.slide_width * scale) / 2
    offset_y = (720 - presentation.slide_height * scale) / 2
    layouts, fonts, warnings = [], set(), []
    for index, slide in enumerate(presentation.slides):
        elements = []
        # Layout/master logos are copied too, while duplicate placeholders use
        # the slide's concrete positions and text.
        shapes = [(True, shape) for shape in slide.slide_layout.slide_master.shapes]
        shapes += [(True, shape) for shape in slide.slide_layout.shapes]
        shapes += [(False, shape) for shape in slide.shapes]
        for number, (inherited, shape) in enumerate(shapes):
            if inherited and shape.is_placeholder:
                continue
            left, top = (
                max(0, offset_x + shape.left * scale),
                max(0, offset_y + shape.top * scale),
            )
            width, height = (
                min(shape.width * scale, 1280 - left),
                min(shape.height * scale, 720 - top),
            )
            if width <= 0 or height <= 0:
                continue
            base = dict(id=f"shape-{number}", x=left, y=top, w=width, h=height)
            if shape.has_table:
                rows = [[cell.text[:300] for cell in row.cells][:12] for row in shape.table.rows][
                    :25
                ]
                elements.append(Element(**base, type="table", rows=rows))
            elif shape.has_chart:
                chart = shape.chart
                try:
                    labels = [str(c.label)[:100] for c in chart.plots[0].categories][:30]
                    values = [float(v or 0) for v in chart.series[0].values][:30]
                    elements.append(Element(**base, type="chart", labels=labels, values=values))
                except (IndexError, ValueError, AttributeError):
                    warnings.append("Unsupported chart omitted")
            elif shape.has_text_frame and shape.text.strip():
                runs = [run for paragraph in shape.text_frame.paragraphs for run in paragraph.runs]
                font = runs[0].font if runs else None
                if font and font.name:
                    fonts.add(font.name)
                size = font.size.pt * 96 / 72 if font and font.size else 24
                elements.append(
                    Element(
                        **base,
                        type="text",
                        text=shape.text[:6000],
                        font_size=max(8, min(150, size)),
                        align={1: "left", 2: "center", 3: "right", 4: "justify"}.get(
                            shape.text_frame.paragraphs[0].alignment, "left"
                        ),
                        valign={1: "top", 2: "middle", 3: "bottom"}.get(
                            shape.text_frame.vertical_anchor, "top"
                        ),
                        color=pptx_color(font.color) if font else "#172033",
                        bold=bool(font.bold) if font else False,
                        max_chars=max(30, min(6000, len(shape.text) * 2)),
                    )
                )
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image = Image.open(io.BytesIO(shape.image.blob)).convert("RGBA")
                output = io.BytesIO()
                image.save(output, "PNG")
                stored = store_asset(
                    asset.owner,
                    f"{job_id}-logo-{index}-{number}.png",
                    output.getvalue(),
                    "image/png",
                    "template-image",
                    idempotency_key=f"{asset.id}:logo:{index}:{number}",
                )
                elements.append(Element(**base, type="image", asset_id=stored.id, editable=False))
            elif shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                try:
                    fill = pptx_color(shape.fill.fore_color, "#FFFFFF")
                except (AttributeError, TypeError):
                    fill = "#FFFFFF"
                elements.append(Element(**base, type="shape", fill=fill, editable=False))
            else:
                warnings.append("Grouped shapes, SmartArt and complex paths require preview review")
        layouts.append(
            Slide(
                id=f"import-{index}",
                name=f"Layout {index + 1}",
                background=pptx_color(slide.background.fill.fore_color, "#FFFFFF")
                if slide.background.fill.type == 1
                else "#FFFFFF",
                elements=elements[:200],
            ).model_dump()
        )
    warnings += [
        f"{font} → Noto Sans CJK SC" for font in sorted(fonts) if font != "Noto Sans CJK SC"
    ]
    warnings.append(
        "Animations and exact master effects are not preserved. Review the imported layouts before saving."
    )
    return {
        "name": Path(asset.name).stem[:120],
        "layouts": layouts,
        "warnings": sorted(set(warnings)),
        "source_asset": asset.id,
        "fonts": sorted(fonts),
    }


def layout_for(template: dict, index: int) -> dict:
    layouts = template["layouts"]
    result = deepcopy(layouts[index % len(layouts)])
    result["id"] = f"slide-{index}"
    return result
