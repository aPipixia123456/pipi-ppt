import io
import json
import zipfile

from pptx import Presentation


def create(client, key="create-one"):
    return client.post(
        "/api/decks",
        headers={"Idempotency-Key": key},
        json={
            "title": "中文产品汇报",
            "topic": "产品发布计划",
            "template_id": "general",
            "slide_count": 5,
            "text_model": "tested-text",
            "images": False,
        },
    )


def test_outline_generate_edit_and_editable_export(site):
    client, calls, _ = site
    from pipi.backend.worker import run_job

    response = create(client)
    assert response.status_code == 200, response.text
    job = response.json()
    assert create(client).json()["id"] == job["id"]
    run_job(job["id"])
    deck = client.get("/api/decks/" + job["deck_id"]).json()
    assert len(deck["outline"]) == 5
    generated = client.post(
        f"/api/decks/{deck['id']}/generate",
        headers={"Idempotency-Key": "generate-one"},
        json={"text_model": "tested-text", "images": False},
    )
    assert generated.status_code == 200, generated.text
    for _ in range(6):
        run_job(generated.json()["id"])
    assert len(calls) == 6  # Duplicate delivery must not call the model again.
    deck = client.get("/api/decks/" + deck["id"]).json()
    assert len(deck["slides"]) == 5
    text = next(e for e in deck["slides"][0]["elements"] if e["type"] == "text")
    text["text"] = "这段文字应当可以在 PowerPoint 和 WPS 中编辑"
    body = {key: deck[key] for key in ("title", "version", "outline", "slides")}
    assert client.put("/api/decks/" + deck["id"], json=body).status_code == 200
    assert client.put("/api/decks/" + deck["id"], json=body).status_code == 409
    export = client.post(
        f"/api/decks/{deck['id']}/export", headers={"Idempotency-Key": "export-one"}
    )
    assert export.status_code == 200, export.text
    run_job(export.json()["id"])
    completed = next(j for j in client.get("/api/jobs").json() if j["id"] == export.json()["id"])
    assert completed["status"] == "complete", completed
    file = client.get("/api/assets/" + completed["result"]["asset_id"])
    assert file.status_code == 200
    presentation = Presentation(io.BytesIO(file.content))
    assert len(presentation.slides) == 5
    assert any("这段文字应当" in s.text for s in presentation.slides[0].shapes if s.has_text_frame)
    assert len(calls) == 6  # Manual editing and export are free of inference.


def test_chart_layout_keeps_text_generation_running(site):
    client, _, _ = site
    from pipi.backend.worker import run_job

    response = client.post(
        "/api/decks",
        headers={"Idempotency-Key": "chart-layout-create"},
        json={
            "title": "财报图表",
            "topic": "季度财报趋势",
            "template_id": "signal",
            "slide_count": 13,
            "text_model": "tested-text",
            "images": False,
        },
    )
    assert response.status_code == 200, response.text
    outline_job = response.json()
    run_job(outline_job["id"])
    deck = client.get("/api/decks/" + outline_job["deck_id"]).json()
    generated = client.post(
        f"/api/decks/{deck['id']}/generate",
        headers={"Idempotency-Key": "chart-layout-generate"},
        json={"text_model": "tested-text", "images": False},
    )
    assert generated.status_code == 200, generated.text
    for _ in range(14):
        run_job(generated.json()["id"])

    job = next(item for item in client.get("/api/jobs").json() if item["id"] == generated.json()["id"])
    assert job["status"] == "complete", job
    assert len(client.get("/api/decks/" + deck["id"]).json()["slides"]) == 13


def test_missing_text_slot_does_not_fail_the_generation(site, monkeypatch):
    client, _, _ = site
    from pipi.backend.worker import run_job

    outline_job = create(client, "missing-slot-create").json()
    run_job(outline_job["id"])
    deck = client.get("/api/decks/" + outline_job["deck_id"]).json()
    generated = client.post(
        f"/api/decks/{deck['id']}/generate",
        headers={"Idempotency-Key": "missing-slot-generate"},
        json={"text_model": "tested-text", "image_model": "", "images": True},
    )
    assert generated.status_code == 200, generated.text
    run_job(generated.json()["id"])  # page-0 uses the normal fake response.

    def malformed_page(*args, **kwargs):
        return {
            "fields": {"example text copied as a key": "Recovered page content"},
            "notes": "Recovered without another model call",
            "image_prompt": "Optional illustration",
        }

    monkeypatch.setattr("pipi.backend.worker.text_json", malformed_page)
    for _ in range(4):
        run_job(generated.json()["id"])

    job = next(item for item in client.get("/api/jobs").json() if item["id"] == generated.json()["id"])
    assert job["status"] == "complete"
    warning_codes = {warning["code"] for warning in job["result"]["warnings"]}
    assert "missing_text_slots" in warning_codes
    assert "image_model_unavailable" in warning_codes
    saved = client.get("/api/decks/" + deck["id"]).json()
    assert len(saved["slides"]) == 5
    text = "\n".join(
        element["text"]
        for slide in saved["slides"]
        for element in slide["elements"]
        if element["type"] == "text"
    )
    assert "Product Overview" not in text
    assert "Our product offers" not in text


def test_outline_response_normalizes_structured_and_short_model_output():
    from pipi.backend.worker import normalize_outline_response

    outline, warnings = normalize_outline_response(
        {
            "outline": [
                {"title": "背景", "key_points": ["现状", "问题"]},
                {"heading": "方案", "summary": "核心做法"},
            ]
        },
        5,
        "新能源项目年度规划",
    )

    assert len(outline) == 5
    assert outline[0] == "背景 — 现状；问题"
    assert outline[1] == "方案 — 核心做法"
    assert "新能源项目年度规划" in outline[-1]
    assert {warning["code"] for warning in warnings} == {"outline_normalized"}


def test_field_values_accepts_list_form():
    from pipi.backend.worker import _field_values

    assert _field_values([{"id": "slot-1", "text": "中文标题"}]) == {
        "slot-1": "中文标题"
    }


def test_text_json_handles_common_model_content_shapes(monkeypatch):
    from pipi.backend.generation import text_json

    responses = iter(
        [
            {
                "choices": [
                    {
                        "message": {
                            "content": [
                                {"type": "text", "text": "```json\n{\"outline\":[\"一页\"]}\n```"}
                            ]
                        }
                    }
                ]
            },
            {"choices": [{"message": {"content": "这不是 JSON"}}]},
        ]
    )
    monkeypatch.setattr(
        "pipi.backend.generation.call_model", lambda *args, **kwargs: next(responses)
    )

    assert text_json("job", "step-1", "model", "prompt") == {"outline": ["一页"]}
    invalid = text_json("job", "step-2", "model", "prompt")
    assert invalid["_pipi_warnings"][0]["code"] == "invalid_model_json"


def test_text_json_enables_kimi_reasoning_without_provider_specific_fields(monkeypatch):
    from pipi.backend.generation import text_json

    captured = {}

    def capture(*args, **kwargs):
        captured.update(args=args, kwargs=kwargs)
        return {"choices": [{"message": {"content": '{"ok":true}'}}]}

    monkeypatch.setattr("pipi.backend.generation.call_model", capture)

    assert text_json("job", "step", "kimi-k3", "prompt") == {"ok": True}
    body = captured["args"][4]
    assert body["reasoning_effort"] == "high"
    assert body["max_tokens"] == 9000
    assert "extra_body" not in body

    text_json("job", "step-off", "kimi-k3", "prompt", reasoning_effort="off")
    off_body = captured["args"][4]
    assert off_body["reasoning_effort"] == "none"


def test_invalid_optional_image_does_not_fail_page_generation(site, monkeypatch):
    client, _, _ = site
    from pipi.backend.worker import run_job

    outline_job = create(client, "invalid-image-create").json()
    run_job(outline_job["id"])
    deck = client.get("/api/decks/" + outline_job["deck_id"]).json()
    generated = client.post(
        f"/api/decks/{deck['id']}/generate",
        headers={"Idempotency-Key": "invalid-image-generate"},
        json={"text_model": "tested-text", "image_model": "tested-image", "images": True},
    )
    assert generated.status_code == 200, generated.text

    def page_with_image(*args, **kwargs):
        prompt = args[3]
        fields = json.loads(
            prompt.split("Template text slots: ", 1)[1].split(". Return JSON", 1)[0]
        )
        return {
            "fields": {slot["id"]: "中文内容" for slot in fields},
            "image_prompt": "Optional illustration",
        }

    monkeypatch.setattr("pipi.backend.worker.text_json", page_with_image)
    monkeypatch.setattr("pipi.backend.worker.call_model", lambda *args, **kwargs: {"data": [{}]})
    for _ in range(5):
        run_job(generated.json()["id"])

    job = next(
        item for item in client.get("/api/jobs").json() if item["id"] == generated.json()["id"]
    )
    assert job["status"] == "complete"
    assert any(warning["code"] == "image_generation_failed" for warning in job["result"]["warnings"])


def test_two_user_isolation_files_jobs_and_csrf(site):
    client, _, _ = site
    first = create(client).json()
    file = client.post(
        "/api/assets?purpose=document",
        files={"file": ("brief.txt", b"private material", "text/plain")},
    ).json()
    client.cookies.set("pipi_session", "browser-two")
    second = create(client).json()
    assert second["id"] != first["id"]
    assert client.get("/api/decks/" + first["deck_id"]).status_code == 404
    assert client.get("/api/assets/" + file["id"]).status_code == 404
    assert client.post("/api/jobs/" + first["id"] + "/cancel").status_code == 404
    client.headers["Origin"] = "https://attacker.example"
    assert client.post("/api/jobs/" + second["id"] + "/cancel").status_code == 403


def test_revocation_blocks_queued_models_and_cookie_access(site):
    client, calls, disabled = site
    job = create(client).json()
    disabled.add(1)
    from pipi.backend.db import Job, session
    from pipi.backend.worker import run_job

    run_job(job["id"])
    with session() as db:
        assert db.get(Job, job["id"]).status == "failed"
    assert not calls
    assert client.get("/api/decks").status_code == 401


def test_upload_limits_archive_and_storage_quota(site):
    client, _, _ = site
    from pipi.backend.db import Account, Policy, session

    assert (
        client.post("/api/assets?purpose=image", files={"file": ("x.svg", b"<svg/>")}).status_code
        == 400
    )
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("ppt/vbaProject.bin", b"macro")
    assert (
        client.post(
            "/api/assets?purpose=template", files={"file": ("x.pptx", output.getvalue())}
        ).status_code
        == 400
    )
    with session() as db:
        policy = db.get(Policy, 1)
        policy.data = {**policy.data, "upload_mb": 1}
        db.get(Account, 1).stored_bytes = policy.data["storage_mb"] * 1024 * 1024
        db.commit()
    assert (
        client.post(
            "/api/assets?purpose=document", files={"file": ("x.txt", b"x" * (1024 * 1024 + 1))}
        ).status_code
        == 413
    )
    assert (
        client.post("/api/assets?purpose=document", files={"file": ("x.txt", b"hello")}).status_code
        == 413
    )


def test_all_builtin_templates_have_editable_layouts(site):
    from pipi.backend.templates import BUILTINS

    client, _, _ = site
    templates = client.get("/api/templates").json()
    assert len(templates) == len(BUILTINS)
    for template in templates:
        assert template["thumbnail"] == f'{template["id"]}/static/thumbnail.png'
        assert client.get("/api/template-assets/" + template["thumbnail"]).status_code == 200
        response = client.get("/api/templates/" + template["id"])
        assert response.status_code == 200, response.text
        assert response.json()["layouts"]


def test_timeout_and_restart_never_replay_paid_step(site, monkeypatch):
    from datetime import timedelta

    import httpx
    from pipi.backend.db import Job, now, session
    from pipi.backend.security import Gateway
    from pipi.backend.worker import dispatch, run_job

    client, calls, _ = site
    job = create(client).json()
    original = Gateway.request

    def timeout(self, method, path, **kwargs):
        if method == "POST":
            calls.append("uncertain")
            raise httpx.ReadTimeout("test timeout")
        return original(self, method, path, **kwargs)

    monkeypatch.setattr(Gateway, "request", timeout)
    run_job(job["id"])
    run_job(job["id"])
    assert calls == ["uncertain"]
    assert client.get("/api/jobs").json()[0]["status"] == "awaiting_confirmation"
    assert client.post("/api/jobs/" + job["id"] + "/resume").status_code == 409
    with session() as db:
        record = db.get(Job, job["id"])
        record.status, record.updated = "running", now() - timedelta(minutes=11)
        db.commit()
    dispatch()
    assert client.get("/api/jobs").json()[0]["status"] == "awaiting_confirmation"
    assert len(calls) == 1
