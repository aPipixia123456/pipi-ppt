import io

import pytest
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


def test_uploaded_template_retains_logo_text_table_and_chart(site):
    client, _, _ = site
    from pipi.backend.conversion import export_deck
    from pipi.backend.db import Asset, session
    from pipi.backend.storage import asset_path
    from pipi.backend.templates import imported_template

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text = slide.shapes.add_textbox(Inches(0.6), Inches(0.6), Inches(8), Inches(1))
    text.text = "中文标题：模板测试"
    text.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    text.text_frame.paragraphs[0].runs[0].font.name = "Missing Brand Font"
    text.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
    picture = io.BytesIO()
    Image.new("RGB", (80, 80), "blue").save(picture, "PNG")
    picture.seek(0)
    slide.shapes.add_picture(picture, Inches(11), Inches(0.4), width=Inches(1))
    table = slide.shapes.add_table(2, 2, Inches(0.6), Inches(2), Inches(5), Inches(2)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "项目", "结果"
    table.cell(1, 0).text, table.cell(1, 1).text = "收入", "120"
    chart = CategoryChartData()
    chart.categories = ["一月", "二月"]
    chart.add_series("收入", [100, 120])
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6), Inches(2), Inches(6), Inches(4), chart
    )
    output = io.BytesIO()
    presentation.save(output)
    upload = client.post(
        "/api/assets?purpose=template", files={"file": ("品牌模板.pptx", output.getvalue())}
    )
    assert upload.status_code == 200, upload.text
    with session() as db:
        source = db.get(Asset, upload.json()["id"])
    template = imported_template(source, "template-test")
    kinds = [element["type"] for element in template["layouts"][0]["elements"]]
    assert all(kind in kinds for kind in ["text", "image", "table", "chart"])
    assert any("Missing Brand Font" in warning for warning in template["warnings"])
    exported = export_deck(1, "品牌模板", template["layouts"], "pptx")
    reopened = Presentation(asset_path(exported))
    shapes = reopened.slides[0].shapes
    assert any(shape.has_chart for shape in shapes)
    assert any(shape.has_table and shape.table.cell(1, 1).text == "120" for shape in shapes)
    assert any(shape.has_text_frame and "中文标题" in shape.text for shape in shapes)
    title = next(shape for shape in shapes if shape.has_text_frame and "中文标题" in shape.text)
    assert title.text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER
    assert any(shape.shape_type == 13 for shape in shapes)


def test_per_user_queue_and_global_claim_limits(site):
    client, _, _ = site
    from pipi.backend.db import Job, Policy, session
    from pipi.backend.worker import claim
    from pipi.tests.test_workflow import create

    jobs = [create(client, f"queue-{i}") for i in range(3)]
    assert all(j.status_code == 200 for j in jobs)
    assert create(client, "over-limit").status_code == 429
    assert claim(jobs[0].json()["id"])
    assert not claim(jobs[1].json()["id"])
    client.cookies.set("pipi_session", "browser-two")
    other = create(client, "another-user").json()
    with session() as db:
        policy = db.get(Policy, 1)
        policy.data = {**policy.data, "generation_concurrency": 1}
        db.commit()
    assert not claim(other["id"])
    with session() as db:
        db.get(Job, jobs[0].json()["id"]).status = "complete"
        db.commit()
    assert claim(other["id"])


def test_same_generated_asset_does_not_double_storage_on_recovery(site):
    from pipi.backend.db import Account, session
    from pipi.backend.storage import store_asset

    one = store_asset(
        1, "test.txt", b"hello", "text/plain", "document", idempotency_key="stable-step"
    )
    two = store_asset(
        1, "test.txt", b"hello", "text/plain", "document", idempotency_key="stable-step"
    )
    assert one.id == two.id
    with session() as db:
        assert db.get(Account, 1).stored_bytes == 5


def test_scene_rejects_infinite_data_and_out_of_canvas_elements():
    from pipi.backend.schema import Element
    from pydantic import ValidationError

    with pytest.raises(ValidationError):
        Element(
            id="chart", type="chart", x=0, y=0, w=100, h=100, labels=["A"], values=[float("nan")]
        )
    with pytest.raises(ValidationError):
        Element(id="text", type="text", x=1200, y=0, w=200, h=100)


def test_generated_long_chinese_text_fits_without_losing_editable_content(site):
    from pipi.backend.conversion import export_deck
    from pipi.backend.schema import Element, Slide
    from pipi.backend.storage import asset_path
    from pipi.backend.worker import fit_text

    text = "这是一段用于确认中文长标题能够完整保留并缩小字号的测试内容"
    element = Element(
        id="title", type="text", x=80, y=80, w=700, h=120, font_size=100, text=text, align="center"
    ).model_dump()
    fit_text(element)
    assert 8 < element["font_size"] < 100
    assert element["text"] == text
    slide = Slide(id="long-text", name="中文长标题", elements=[element]).model_dump()
    exported = export_deck(1, "长标题", [slide], "pptx")
    result = Presentation(asset_path(exported)).slides[0].shapes[0]
    assert result.text == text
    assert result.text_frame.paragraphs[0].runs[0].font.size.pt == pytest.approx(
        element["font_size"] * 0.75
    )
